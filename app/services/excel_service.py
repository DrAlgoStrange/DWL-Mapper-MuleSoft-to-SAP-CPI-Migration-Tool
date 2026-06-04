"""
Excel Export Service — generates the CPI Mapping Sheet .xlsx file
modelled after the CPI_SalesOrder_Mapping_LegendToE2.xlsx reference.
"""
import os
import logging
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# ── Colour Palette ──────────────────────────────────────────────────────────
DARK_HEADER   = "1A1A2E"   # deep navy
MID_HEADER    = "16213E"   # section header
ACCENT        = "0F3460"   # sub-header / segment rows
LIGHT_BG      = "F8F9FA"   # alternate row light
WHITE         = "FFFFFF"
YELLOW_BG     = "FFF9C4"   # hardcoded
GREEN_BG      = "E8F5E9"   # pass-through
ORANGE_BG     = "FFF3E0"   # transformation
BLUE_BG       = "E3F2FD"   # config property
PURPLE_BG     = "F3E5F5"   # conditional
SEGMENT_BG    = "E8EAF6"   # segment section header row

TEXT_LIGHT    = "FFFFFF"
TEXT_DARK     = "1A1A2E"
TEXT_ACCENT   = "0F3460"

MAPPING_TYPE_STYLES = {
    "passthrough":    {"bg": GREEN_BG,  "emoji": "🟢"},
    "hardcoded":      {"bg": YELLOW_BG, "emoji": "🟡"},
    "transformation": {"bg": ORANGE_BG, "emoji": "🟠"},
    "config":         {"bg": BLUE_BG,   "emoji": "⚙️"},
    "conditional":    {"bg": PURPLE_BG, "emoji": "⚠️"},
}

THIN = Side(style='thin', color="CCCCCC")
MED  = Side(style='medium', color="999999")

def _thin_border():
    return Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def _med_border():
    return Border(left=MED, right=MED, top=MED, bottom=MED)

def _hfill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, color=TEXT_DARK, size=10, italic=False):
    return Font(name="Arial", bold=bold, color=color, size=size, italic=italic)

def _align(h="left", v="center", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def generate_mapping_xlsx(mapping_data: dict, output_path: str, project_name: str) -> str:
    """
    Generate a professional CPI Mapping Sheet .xlsx file.
    Returns the output_path on success.
    """
    try:
        wb = Workbook()

        # ── Sheet 1: CPI Mapping Sheet ─────────────────────────────────────
        ws = wb.active
        ws.title = "CPI Mapping Sheet"
        _build_mapping_sheet(ws, mapping_data, project_name)

        # ── Sheet 2: Config Parameters ─────────────────────────────────────
        ws2 = wb.create_sheet("Config Parameters")
        _build_config_sheet(ws2, mapping_data)

        # ── Sheet 3: CPI Notes ─────────────────────────────────────────────
        ws3 = wb.create_sheet("CPI Implementation Notes")
        _build_notes_sheet(ws3, mapping_data, project_name)

        wb.save(output_path)
        logger.info(f"Mapping XLSX saved: {output_path}")
        return output_path

    except Exception as e:
        logger.error(f"Error generating XLSX: {e}", exc_info=True)
        raise


def _build_mapping_sheet(ws, mapping_data: dict, project_name: str):
    """Build the main CPI Mapping Sheet tab."""

    # ── Title block ────────────────────────────────────────────────────────
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = f"SAP CPI Message Mapping — {mapping_data.get('source_system', 'Source')} → {mapping_data.get('target_system', 'Target')}"
    title_cell.font = _font(bold=True, color=TEXT_LIGHT, size=14)
    title_cell.fill = _hfill(DARK_HEADER)
    title_cell.alignment = _align("center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:I2")
    proj_cell = ws["A2"]
    proj_cell.value = f"Project: {project_name}  |  Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}  |  {mapping_data.get('summary', '')}"
    proj_cell.font = _font(italic=True, color=TEXT_LIGHT, size=9)
    proj_cell.fill = _hfill(MID_HEADER)
    proj_cell.alignment = _align("center")
    ws.row_dimensions[2].height = 20

    # ── Legend ─────────────────────────────────────────────────────────────
    ws.merge_cells("A3:I3")
    legend_cell = ws["A3"]
    legend_cell.value = "LEGEND:  🟡 Fixed/Hardcoded Value   🟢 Direct Pass-Through (1:1 copy)   🟠 Transformation / Logic   ⚙️ Config Property   ⚠️ Conditional / Complex Logic"
    legend_cell.font = _font(bold=True, size=9, color=TEXT_DARK)
    legend_cell.fill = _hfill("F0F0F0")
    legend_cell.alignment = _align("center", wrap=False)
    ws.row_dimensions[3].height = 18

    # ── Column Headers ─────────────────────────────────────────────────────
    headers = [
        "IDoc Segment / Element",
        "Target Field",
        "Field Description",
        "Type",
        "Len",
        "Source Field",
        "Mapping Logic / Value",
        "Config Property",
        "Notes / SAP CPI Guidance",
    ]
    col_widths = [28, 22, 28, 8, 6, 32, 40, 30, 55]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = _font(bold=True, color=TEXT_LIGHT, size=10)
        cell.fill = _hfill(ACCENT)
        cell.alignment = _align("center")
        cell.border = _med_border()
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[4].height = 22
    ws.freeze_panes = "A5"

    # ── Data Rows ──────────────────────────────────────────────────────────
    current_row = 5
    mapping_rows = mapping_data.get("mapping_rows", [])
    last_segment = None

    for row_data in mapping_rows:
        segment = row_data.get("segment", "")

        # Insert segment section header when segment changes
        if segment and segment != last_segment:
            ws.merge_cells(f"A{current_row}:I{current_row}")
            seg_cell = ws.cell(row=current_row, column=1, value=f"▶  {segment}")
            seg_cell.font = _font(bold=True, color=ACCENT, size=10)
            seg_cell.fill = _hfill(SEGMENT_BG)
            seg_cell.alignment = _align("left", wrap=False)
            seg_cell.border = _med_border()
            ws.row_dimensions[current_row].height = 20
            current_row += 1
            last_segment = segment

        mapping_type = row_data.get("mapping_type", "passthrough").lower()
        style = MAPPING_TYPE_STYLES.get(mapping_type, MAPPING_TYPE_STYLES["passthrough"])
        row_bg = style["bg"]
        emoji   = style["emoji"]

        mapping_logic_raw = row_data.get("mapping_logic", "")
        mapping_logic_display = f"{emoji}  {mapping_logic_raw}" if not mapping_logic_raw.startswith(("🟢","🟡","🟠","⚙","⚠")) else mapping_logic_raw

        row_values = [
            segment,
            row_data.get("target_field", ""),
            row_data.get("field_description", ""),
            row_data.get("type", ""),
            row_data.get("length", ""),
            row_data.get("source_field", ""),
            mapping_logic_display,
            row_data.get("config_property", ""),
            row_data.get("notes", ""),
        ]

        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = _font(size=9, color=TEXT_DARK)
            cell.fill = _hfill(row_bg)
            cell.alignment = _align("left")
            cell.border = _thin_border()

        ws.row_dimensions[current_row].height = 36
        current_row += 1

    # ── Totals footer ─────────────────────────────────────────────────────
    ws.merge_cells(f"A{current_row}:I{current_row}")
    total_cell = ws.cell(row=current_row, column=1,
        value=f"Total Mapped Fields: {len(mapping_rows)}")
    total_cell.font = _font(bold=True, color=TEXT_LIGHT, size=10)
    total_cell.fill = _hfill(DARK_HEADER)
    total_cell.alignment = _align("right")

    # ── Print / view settings ──────────────────────────────────────────────
    ws.sheet_view.showGridLines = True
    ws.print_title_rows = "1:4"


def _build_config_sheet(ws, mapping_data: dict):
    """Build the Config Parameters sheet."""
    ws.merge_cells("A1:D1")
    h = ws["A1"]
    h.value = "⚙️  CPI Configuration Parameters — to be set as Integration Flow Parameters / Externalized Config"
    h.font = _font(bold=True, color=TEXT_LIGHT, size=12)
    h.fill = _hfill(DARK_HEADER)
    h.alignment = _align("center")
    ws.row_dimensions[1].height = 26

    headers = ["MuleSoft Property Key", "Sample Value", "Target IDoc Field", "SAP CPI Guidance"]
    col_widths = [35, 30, 30, 60]
    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = _font(bold=True, color=TEXT_LIGHT, size=10)
        cell.fill = _hfill(ACCENT)
        cell.alignment = _align("center")
        cell.border = _med_border()
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[2].height = 20
    ws.freeze_panes = "A3"

    config_params = mapping_data.get("config_parameters", [])
    for row_idx, param in enumerate(config_params, start=3):
        row_bg = BLUE_BG if row_idx % 2 == 0 else WHITE
        row_vals = [
            param.get("mulesoft_key", ""),
            param.get("sample_value", ""),
            param.get("target_field", ""),
            param.get("cpi_guidance", ""),
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = _font(size=9)
            cell.fill = _hfill(row_bg)
            cell.alignment = _align("left")
            cell.border = _thin_border()
        ws.row_dimensions[row_idx].height = 28

    if not config_params:
        ws.merge_cells("A3:D3")
        ws["A3"].value = "No config properties identified."
        ws["A3"].font = _font(italic=True, color="888888")
        ws["A3"].alignment = _align("center")


def _build_notes_sheet(ws, mapping_data: dict, project_name: str):
    """Build the CPI Implementation Notes sheet."""
    ws.merge_cells("A1:B1")
    h = ws["A1"]
    h.value = f"CPI Implementation Notes — {project_name}"
    h.font = _font(bold=True, color=TEXT_LIGHT, size=12)
    h.fill = _hfill(DARK_HEADER)
    h.alignment = _align("center")
    ws.row_dimensions[1].height = 26

    headers = ["#", "Implementation Note"]
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 110
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font = _font(bold=True, color=TEXT_LIGHT, size=10)
        cell.fill = _hfill(ACCENT)
        cell.alignment = _align("center")
        cell.border = _med_border()
    ws.row_dimensions[2].height = 20
    ws.freeze_panes = "A3"

    notes = mapping_data.get("cpi_implementation_notes", [])
    for row_idx, note in enumerate(notes, start=3):
        row_bg = LIGHT_BG if row_idx % 2 == 0 else WHITE
        ws.cell(row=row_idx, column=1, value=row_idx - 2).font = _font(bold=True, size=9)
        ws.cell(row=row_idx, column=1).fill = _hfill(row_bg)
        ws.cell(row=row_idx, column=1).alignment = _align("center")
        ws.cell(row=row_idx, column=1).border = _thin_border()

        note_cell = ws.cell(row=row_idx, column=2, value=note)
        note_cell.font = _font(size=9)
        note_cell.fill = _hfill(row_bg)
        note_cell.alignment = _align("left")
        note_cell.border = _thin_border()
        ws.row_dimensions[row_idx].height = 30

    if not notes:
        ws.merge_cells("A3:B3")
        ws["A3"].value = "No additional notes."
        ws["A3"].font = _font(italic=True, color="888888")
        ws["A3"].alignment = _align("center")
